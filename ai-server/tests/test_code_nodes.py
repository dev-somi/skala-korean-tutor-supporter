"""
이 파일의 역할: LLM을 쓰지 않는 코드(로드·구성안 노드, 수업 틀, PPT 모듈)를 테스트한다.
PPT 모듈은 그래프에서 뗐지만(ADR-0015) 나중에 다시 쓸 수 있게 모듈 단위로 시험해 둔다.

실행 방법: uv run pytest
"""

import pytest
from pptx import Presentation

from app.agent.lesson_prep import lesson_frame, nodes, pptx_builder, render
from app.agent.lesson_prep.schemas import CATEGORIES


def make_state(tmp_path, monkeypatch, explanation="説明です。") -> dict:
    """PPT를 만들 수 있을 만큼 채워진 State를 손으로 만든다."""
    monkeypatch.setattr(pptx_builder, "OUTPUT_DIR", tmp_path)
    state = {"student_id": "student_sample", "unit_no": None, "lesson_minutes": 50, "target_lang": "ja"}
    state.update(nodes.load_profile(state))
    state.update(nodes.load_textbook(state))
    item_ids = {"review": ["u02-g01"], "new_vocab": ["u03-v01"], "new_grammar": ["u03-g01"], "homework": ["u03-e01"]}
    minutes = {"review": 10, "new_vocab": 15, "new_grammar": 20, "homework": 5}
    state["drafts"] = {}
    for category in CATEGORIES:
        state["drafts"][category] = {
            "item_ids": item_ids[category],
            "minutes": minutes[category],
            "reason": "테스트",
            "explanation": explanation,
            "examples": [{"korean": "새로 만든 예문입니다.", "is_new": True}],
            "script": f"{category}の台本です。",
        }
    return state


def all_slide_text(pptx_path: str) -> str:
    """PPT의 모든 슬라이드 글과 노트를 한 덩어리 글로 모은다."""
    texts = []
    for slide in Presentation(pptx_path).slides:
        for shape in slide.shapes:
            if shape.has_text_frame:
                texts.append(shape.text_frame.text)
        if slide.has_notes_slide:
            texts.append(slide.notes_slide.notes_text_frame.text)
    return "\n".join(texts)


def test_load_nodes():
    """프로필의 next_unit(3과)을 읽고, 과거에 배운 2과 항목을 복습 후보로 모은다."""
    state = {"student_id": "student_sample", "unit_no": None}
    state.update(nodes.load_profile(state))
    state.update(nodes.load_textbook(state))

    assert state["profile"]["native_language"] == "ja"
    assert state["textbook_unit"]["unit_no"] == 3
    past_ids = {item["id"] for item in state["past_items"]}
    assert past_ids == nodes.learned_ids(state["profile"])


def test_unit_no_overrides_next_unit():
    """선생님이 과 번호를 직접 주면 프로필의 next_unit 대신 그 과를 읽는다."""
    state = {"student_id": "student_sample", "unit_no": 2}
    state.update(nodes.load_profile(state))
    state.update(nodes.load_textbook(state))
    assert state["textbook_unit"]["unit_no"] == 2


def test_build_pptx_keeps_textbook_text(tmp_path, monkeypatch):
    """PPT에 교재 원문, 설명, 새 예문 표시, 대본(노트)이 들어간다."""
    state = make_state(tmp_path, monkeypatch)
    path = pptx_builder.build_pptx_file(state)
    text = all_slide_text(path)

    assert "노래 (명사) : song" in text  # 교재 원문 그대로
    assert "説明です。" in text  # 개인화 설명
    assert "새로 만든 예문입니다. (새 예문)" in text  # 새 예문 표시
    assert "new_grammarの台本です。" in text  # 대본은 노트에
    assert render.check_overflow(path) == []  # 넘침 없음


def test_long_text_is_split_into_slides(tmp_path, monkeypatch):
    """설명이 길면 슬라이드를 여러 장으로 나눠서 넘치지 않게 한다."""
    long_explanation = "\n".join(f"{i}番目の説明です。" for i in range(60))
    short = make_state(tmp_path, monkeypatch)
    short_count = len(Presentation(pptx_builder.build_pptx_file(short)).slides)

    long = make_state(tmp_path, monkeypatch, explanation=long_explanation)
    path = pptx_builder.build_pptx_file(long)

    assert len(Presentation(path).slides) > short_count
    assert render.check_overflow(path) == []


def test_overflow_is_detected_and_smaller_font_fixes_it(tmp_path, monkeypatch):
    """줄바꿈 없는 아주 긴 문단은 넘침으로 잡히고, 글자를 줄여 다시 만들면 해결된다."""
    monkeypatch.setattr(render, "find_soffice", lambda: None)
    state = make_state(tmp_path, monkeypatch, explanation="あ" * 420)

    assert render.check_render(pptx_builder.build_pptx_file(state)) != []

    # render_retry가 올라가면 글자를 줄여서 만든다
    state["render_retry"] = 1
    assert render.check_render(pptx_builder.build_pptx_file(state)) == []


@pytest.mark.skipif(render.find_soffice() is None, reason="LibreOffice가 설치되어 있지 않음")
def test_pdf_conversion(tmp_path, monkeypatch):
    """LibreOffice로 PDF 변환이 되고, 쪽수가 슬라이드 수와 같다."""
    state = make_state(tmp_path, monkeypatch)
    path = pptx_builder.build_pptx_file(state)
    assert render.check_render(path) == []


def plan_state(lesson_minutes=50, student_id="student_sample") -> dict:
    state = {"student_id": student_id, "unit_no": None, "lesson_minutes": lesson_minutes}
    state.update(nodes.load_profile(state))
    state.update(nodes.load_textbook(state))
    return state


def test_plan_lesson_follows_rules():
    """약점 → 직전 수업 순서로 복습, 이번 과 어휘·문법은 교재 순서 그대로, 숙제는 문법 수만큼 연습문제."""
    state = plan_state()
    plan = nodes.plan_lesson(state)["plan"]
    unit = state["textbook_unit"]

    assert plan["review"]["item_ids"][:3] == ["u02-g01", "u02-g02", "u02-v05"]  # 프로필의 약점이 먼저
    assert len(plan["review"]["item_ids"]) <= nodes.MAX_REVIEW_ITEMS
    assert plan["new_vocab"]["item_ids"] == [v["id"] for v in unit["vocabulary"]]
    assert plan["new_grammar"]["item_ids"] == [g["id"] for g in unit["grammar"]]
    assert len(plan["homework"]["item_ids"]) == len(unit["grammar"])


@pytest.mark.parametrize("lesson_minutes", [30, 40, 50, 60, 90])
def test_frame_minutes_always_add_up(lesson_minutes):
    """어떤 수업 길이든 7단계 시간 합계 = 수업 시간, 범주 시간 + 틀 전용 단계 = 수업 시간."""
    for has_review in (True, False):
        steps = lesson_frame.step_minutes(lesson_minutes, has_review)
        assert sum(steps.values()) == lesson_minutes
        cats = lesson_frame.category_minutes(steps)
        assert sum(cats.values()) + lesson_frame.frame_only_minutes(steps) == lesson_minutes
    assert lesson_frame.step_minutes(lesson_minutes, has_review=False)["review"] == 0


def test_assemble_fills_frame():
    """틀의 7단계 순서는 고정이고, 워밍업·활용 활동에는 브리핑 내용이 들어간다."""
    drafts = {"review": {"item_ids": ["u02-g01"]}}
    briefing = {"warmup_questions": ["주말에 뭐 했어요?"], "practice_activity": "카페 주문 롤플레이"}
    frame = lesson_frame.assemble(50, drafts, briefing)

    assert [s["key"] for s in frame] == ["warmup", "review", "present", "drill", "activity", "feedback", "wrapup"]
    assert sum(s["minutes"] for s in frame) == 50
    by_key = {s["key"]: s for s in frame}
    assert by_key["warmup"]["items"] == ["주말에 뭐 했어요?"]
    assert by_key["activity"]["items"] == ["카페 주문 롤플레이"]
    assert by_key["present"]["categories"] == ["new_vocab", "new_grammar"]


def test_weakness_without_history_is_reviewed():
    """수업 기록이 없어도 앞 과의 약점은 복습에 들어간다 (체험 수업 뒤 바로 등록한 학생)."""
    state = {"student_id": "student_sample", "unit_no": None, "lesson_minutes": 50}
    state.update(nodes.load_profile(state))
    state["profile"]["lesson_history"] = []
    state.update(nodes.load_textbook(state))
    plan = nodes.plan_lesson(state)["plan"]

    assert plan["review"]["item_ids"] == ["u02-g01", "u02-g02", "u02-v05"]
    assert plan["review"]["minutes"] > 0
