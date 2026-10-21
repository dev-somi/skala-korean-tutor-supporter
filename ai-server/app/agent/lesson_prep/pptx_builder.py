"""
이 파일의 역할: 승인된 수업 자료(State)를 PPT 파일로 만든다. LLM은 쓰지 않는다.

슬라이드 구성
  1. 표지: 과 제목, 학생, 범주별 시간 배분
  2. 범주마다
     - "교재 내용" 슬라이드: 교재 항목 원문 (교재 JSON에서 그대로 가져옴)
     - "설명" 슬라이드: 개인화 설명 + 예문
     - 대본은 슬라이드 노트에 넣는다 (선생님만 보는 내용)

내용이 길면 여러 장으로 나눈다.
"""

import math
import sys
from pathlib import Path

from pptx import Presentation
from pptx.oxml.ns import qn
from pptx.util import Inches, Pt

from app.agent.lesson_prep.schemas import CATEGORIES, CATEGORY_NAMES
from app.agent.lesson_prep.textbook import PROJECT_DIR, all_items_by_id, format_item

TEMPLATE_PATH = PROJECT_DIR / "templates" / "lesson_template.pptx"
OUTPUT_DIR = PROJECT_DIR / "output"

# --- 슬라이드 위 글상자의 위치와 크기 (단위: 인치) ---
BOX_LEFT = 0.6
BOX_WIDTH = 12.1
TITLE_TOP = 0.4
TITLE_HEIGHT = 1.0
BODY_TOP = 1.6
BODY_HEIGHT = 5.4

# 한글과 일본어(가나, 한자)가 모두 들어 있는 글꼴. 글꼴을 지정하지 않으면
# PDF로 바꿀 때 글자가 빠질 수 있다. 윈도우에서는 "Malgun Gothic" 등으로 바꾼다.
FONT_NAME = "Arial Unicode MS"

TITLE_FONT = 32  # 제목 글자 크기(pt)
BASE_BODY_FONT = 24  # 본문 기본 글자 크기(pt)
FONT_STEP = 3  # 렌더 점검에 실패할 때마다 줄이는 크기(pt)
LINE_SPACING = 1.25  # 줄 높이 = 글자 크기 x 1.25 (대략적인 값)
SAFETY = 0.9  # 글상자 안쪽 여백을 고려해 실제 공간의 90%만 쓴다고 가정


# ---------------------------------------------------------------------------
# 글자가 차지하는 공간을 어림잡는 함수들 (render.py의 넘침 점검에서도 사용)
# ---------------------------------------------------------------------------


def text_width(text: str) -> float:
    """글의 가로 길이를 '글자 크기의 몇 배인지'로 어림잡는다."""
    width = 0.0
    for ch in text:
        # 한글/한자/가나 같은 전각 글자는 1칸, 영어/숫자/공백은 약 0.55칸
        if ord(ch) > 0x2E80:
            width += 1.0
        else:
            width += 0.55
    return width


def count_wrapped_lines(text: str, box_width_inch: float, font_pt: float) -> int:
    """한 문단이 글상자 안에서 자동 줄바꿈되면 몇 줄이 되는지 어림잡는다."""
    # 1인치 = 72pt. 한 줄에 들어가는 글자 수 = 글상자 가로 / 글자 크기
    chars_per_line = box_width_inch * 72 * SAFETY / font_pt
    return max(1, math.ceil(text_width(text) / chars_per_line))


def max_lines(box_height_inch: float, font_pt: float) -> int:
    """글상자에 세로로 최대 몇 줄이 들어가는지 어림잡는다."""
    return int(box_height_inch * 72 * SAFETY / (font_pt * LINE_SPACING))


def split_into_pages(lines: list[str], font_pt: float) -> list[list[str]]:
    """줄 목록을 슬라이드 한 장에 들어갈 만큼씩 나눈다."""
    limit = max_lines(BODY_HEIGHT, font_pt)

    # 1) 줄들을 "묶음"으로 나눈다. 들여쓴 줄(예문 등)은 바로 앞 줄과 같은 묶음이다.
    #    이렇게 하면 단어와 그 예문이 서로 다른 슬라이드로 갈라지지 않는다.
    blocks = []
    for line in lines:
        if line.startswith(" ") and blocks:
            blocks[-1].append(line)
        else:
            blocks.append([line])

    # 2) 묶음 단위로 슬라이드를 채운다.
    pages = []
    current = []  # 지금 채우고 있는 슬라이드의 줄들
    used = 0  # 지금 슬라이드에서 이미 쓴 줄 수
    for block in blocks:
        need = sum(count_wrapped_lines(line, BOX_WIDTH, font_pt) for line in block)
        # 이 묶음을 넣으면 넘치는 경우, 지금까지 모은 것을 한 장으로 확정하고 새 장을 시작
        if current and used + need > limit:
            pages.append(current)
            current = []
            used = 0
        current.extend(block)
        used += need
    if current:
        pages.append(current)
    return pages


# ---------------------------------------------------------------------------
# 슬라이드를 만드는 함수들
# ---------------------------------------------------------------------------


def set_font(run, font_pt: float, bold: bool) -> None:
    """글자 조각(run)의 글꼴, 크기, 굵기를 정한다."""
    run.font.size = Pt(font_pt)
    run.font.bold = bold
    run.font.name = FONT_NAME  # 영어/숫자용 글꼴
    # 한글/일본어용(동아시아) 글꼴은 python-pptx에 전용 기능이 없어서 XML에 직접 적는다
    properties = run._r.get_or_add_rPr()  # 이 글자 조각의 서식이 적힌 XML 부분
    east_asian = properties.makeelement(qn("a:ea"), {"typeface": FONT_NAME})
    properties.append(east_asian)


def add_textbox(slide, top: float, height: float, lines: list[str], font_pt: float, bold: bool = False):
    """슬라이드에 글상자 하나를 추가하고, 줄마다 문단 하나씩 넣는다."""
    box = slide.shapes.add_textbox(Inches(BOX_LEFT), Inches(top), Inches(BOX_WIDTH), Inches(height))
    frame = box.text_frame
    frame.word_wrap = True  # 글상자 가로를 넘으면 자동 줄바꿈
    for i, line in enumerate(lines):
        # 글상자에는 빈 문단이 처음부터 1개 있으므로 첫 줄은 그것을 쓴다
        paragraph = frame.paragraphs[0] if i == 0 else frame.add_paragraph()
        run = paragraph.add_run()
        run.text = line
        set_font(run, font_pt, bold)


def add_slide(prs, title: str, body_lines: list[str], body_font: float, notes: str = ""):
    """제목 + 본문 + (선택) 노트가 있는 슬라이드 한 장을 추가한다."""
    blank_layout = prs.slide_layouts[6]  # 6번 = 아무것도 없는 빈 레이아웃
    slide = prs.slides.add_slide(blank_layout)
    add_textbox(slide, TITLE_TOP, TITLE_HEIGHT, [title], TITLE_FONT, bold=True)
    add_textbox(slide, BODY_TOP, BODY_HEIGHT, body_lines, body_font)
    if notes:
        slide.notes_slide.notes_text_frame.text = notes
    return slide


def item_lines(item_ids: list[str], items: dict) -> list[str]:
    """항목 ID 목록을 교재 원문 줄 목록으로 바꾼다. 교재에 없는 ID는 표시만 해 둔다."""
    lines = []
    for item_id in item_ids:
        if item_id in items:
            lines.extend(format_item(items[item_id]))
        else:
            lines.append(f"(교재에서 찾을 수 없는 항목: {item_id})")
    return lines


def explanation_lines(draft: dict) -> list[str]:
    """개인화 설명과 예문을 줄 목록으로 바꾼다."""
    # 설명 글은 줄바꿈 기준으로 나누고, 빈 줄은 뺀다
    lines = [line for line in draft["explanation"].split("\n") if line.strip()]
    for example in draft.get("examples", []):
        # 교재에 없는 새 예문은 눈에 보이게 표시한다
        mark = " (새 예문)" if example["is_new"] else ""
        lines.append(f"   예) {example['korean']}{mark}")
    return lines


def build_pptx_file(state: dict) -> str:
    """State의 내용으로 .pptx 파일을 만들고 파일 경로를 돌려준다."""
    # 템플릿이 없으면 먼저 만든다
    if not TEMPLATE_PATH.exists():
        sys.path.insert(0, str(TEMPLATE_PATH.parent))
        from make_template import make_template

        make_template()

    prs = Presentation(str(TEMPLATE_PATH))
    unit = state["textbook_unit"]
    profile = state["profile"]
    drafts = state["drafts"]
    items = all_items_by_id(state)

    # 렌더 점검에 실패해서 다시 만드는 경우, 실패한 횟수만큼 글자를 줄인다
    body_font = BASE_BODY_FONT - FONT_STEP * state.get("render_retry", 0)

    # --- 1. 표지 ---
    cover_lines = [
        f"학생: {profile['name']}",
        f"수업 시간: {state['lesson_minutes']}분",
        "",
    ]
    for category in CATEGORIES:
        cover_lines.append(f"{CATEGORY_NAMES[category]}: {drafts[category]['minutes']}분")
    add_slide(prs, f"{unit['unit_no']}과 {unit['title']}", cover_lines, body_font)

    # --- 2. 범주별 슬라이드 ---
    for category in CATEGORIES:
        draft = drafts[category]
        name = f"{CATEGORY_NAMES[category]} ({draft['minutes']}분)"
        script = draft["script"]  # 대본은 이 범주의 모든 슬라이드 노트에 넣는다

        # 교재 내용 슬라이드 (길면 여러 장)
        for page in split_into_pages(item_lines(draft["item_ids"], items), body_font):
            add_slide(prs, f"{name} - 교재 내용", page, body_font, notes=script)

        # 설명 슬라이드 (길면 여러 장)
        for page in split_into_pages(explanation_lines(draft), body_font):
            add_slide(prs, f"{name} - 설명", page, body_font, notes=script)

    # --- 3. 저장 ---
    # API로 실행하면 output_subdir(예: "<student_id>/<lesson_id>")에 학생·수업별로 저장한다
    out_dir = OUTPUT_DIR / state["output_subdir"] if state.get("output_subdir") else OUTPUT_DIR
    out_dir.mkdir(parents=True, exist_ok=True)
    file_name = f"lesson_unit{unit['unit_no']:02d}_{profile['student_id']}_{state['target_lang']}.pptx"
    path = out_dir / file_name
    prs.save(str(path))
    return str(path)
