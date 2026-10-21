"""
이 파일의 역할: 만들어진 .pptx에 글자 넘침이나 깨짐이 없는지 점검한다. LLM은 쓰지 않는다.

점검 2가지
  1. 글자 넘침: 글상자 크기와 글자 수로, 글이 글상자 밖으로 나가는지 어림잡는다.
  2. PDF 변환: LibreOffice가 설치되어 있으면 PDF로 바꿔 보고,
     변환이 되는지 + PDF 쪽수가 슬라이드 수와 같은지 + 글꼴이 바뀌지 않았는지 확인한다.
     (LibreOffice가 없으면 2번은 건너뛴다)
"""

import shutil
import subprocess
from pathlib import Path

from pptx import Presentation
from pypdf import PdfReader

from app.agent.lesson_prep.pptx_builder import FONT_NAME, LINE_SPACING, SAFETY, count_wrapped_lines

# macOS에서 LibreOffice가 설치되는 기본 위치
MAC_SOFFICE = "/Applications/LibreOffice.app/Contents/MacOS/soffice"


def find_soffice() -> str | None:
    """LibreOffice 실행 파일(soffice)의 경로를 찾는다. 없으면 None."""
    found = shutil.which("soffice") or shutil.which("libreoffice")
    if found:
        return found
    if Path(MAC_SOFFICE).exists():
        return MAC_SOFFICE
    return None


def convert_to_pdf(pptx_path: str) -> str | None:
    """.pptx를 같은 폴더에 PDF로 변환한다. 성공하면 PDF 경로, 실패하면 None."""
    soffice = find_soffice()
    if soffice is None:
        return None
    out_dir = str(Path(pptx_path).parent)
    # --convert-to: 창을 띄우지 않고 변환만 한다.
    # 주의: macOS에서 --headless 옵션을 함께 주면 시스템 글꼴을 찾지 못해
    #       한글/일본어 글자가 PDF에서 사라지므로 넣지 않는다.
    command = [soffice, "--convert-to", "pdf", "--outdir", out_dir, pptx_path]
    try:
        subprocess.run(command, capture_output=True, timeout=120)
    except subprocess.TimeoutExpired:
        return None
    pdf_path = Path(pptx_path).with_suffix(".pdf")
    return str(pdf_path) if pdf_path.exists() else None


def check_overflow(pptx_path: str) -> list[str]:
    """글상자마다 글이 넘치는지 어림잡아, 문제 설명 목록을 돌려준다. (없으면 빈 목록)"""
    problems = []
    prs = Presentation(pptx_path)
    for slide_no, slide in enumerate(prs.slides, start=1):
        for shape in slide.shapes:
            if not shape.has_text_frame:
                continue
            # EMU 단위를 인치로 바꾼다 (1인치 = 914400 EMU)
            width_inch = shape.width / 914400
            height_inch = shape.height / 914400

            # 문단마다 필요한 높이(pt)를 더한다
            needed_pt = 0.0
            for paragraph in shape.text_frame.paragraphs:
                if not paragraph.runs:
                    continue
                font_pt = paragraph.runs[0].font.size.pt
                lines = count_wrapped_lines(paragraph.text, width_inch, font_pt)
                needed_pt += lines * font_pt * LINE_SPACING

            available_pt = height_inch * 72 * SAFETY
            if needed_pt > available_pt:
                problems.append(
                    f"{slide_no}번 슬라이드: 글이 글상자를 넘칩니다 "
                    f"(필요 {needed_pt:.0f}pt > 가능 {available_pt:.0f}pt)"
                )
    return problems


def check_render(pptx_path: str) -> list[str]:
    """넘침 점검 + PDF 변환 점검을 모두 하고, 문제 설명 목록을 돌려준다."""
    problems = check_overflow(pptx_path)

    # LibreOffice가 없으면 PDF 점검은 건너뛴다
    if find_soffice() is None:
        print("[render_check] LibreOffice가 없어 PDF 변환 점검은 건너뜁니다.")
        return problems

    pdf_path = convert_to_pdf(pptx_path)
    if pdf_path is None:
        problems.append("PDF로 변환하지 못했습니다 (파일이 깨졌을 수 있음)")
        return problems

    # PDF 쪽수와 슬라이드 수가 다르면 변환 중에 문제가 생긴 것이다
    slide_count = len(Presentation(pptx_path).slides)
    reader = PdfReader(pdf_path)
    if len(reader.pages) != slide_count:
        problems.append(f"PDF 쪽수({len(reader.pages)})가 슬라이드 수({slide_count})와 다릅니다")

    # 글자 깨짐 점검: 우리가 지정한 글꼴이 PDF에 실제로 쓰였는지 본다.
    # 그 글꼴을 찾지 못하면 다른 글꼴로 바뀌고, 한글/일본어 글자가 빈칸으로 나온다.
    wanted = FONT_NAME.replace(" ", "")  # PDF 안에서는 글꼴 이름에 공백이 없다
    used = pdf_font_names(reader)
    if not any(wanted in name for name in used):
        problems.append(f"글자가 깨졌을 수 있습니다: '{FONT_NAME}' 글꼴 대신 {sorted(used)} 글꼴로 바뀌었습니다")
    return problems


def pdf_font_names(reader: PdfReader) -> set:
    """PDF의 모든 쪽에서 쓰인 글꼴 이름을 모은다."""
    names = set()
    for page in reader.pages:
        fonts = page.get("/Resources", {}).get("/Font", {})
        for font in fonts.values():
            names.add(str(font.get_object().get("/BaseFont", "")))
    return names
