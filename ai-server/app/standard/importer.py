"""
이 파일의 역할: 국립국어원 「국제 통용 한국어 표준 교육과정 적용 연구(4단계)」(2017)의
어휘·문법 등급 목록 엑셀을 서비스가 읽는 JSON으로 바꾼다. (기술 ADR-0016)

  data/standard/source/*.xlsx  원본 (공공누리 제1유형, 출처표시)
  data/standard/grammar.json   문법·표현 336개 (1~6급)
  data/standard/vocab.json     어휘 10,635개 (1~6급)

ID는 엑셀의 "전체 번호"로 만든다 (std-g-007, std-v-00001). 원본을 다시 받아도 번호가 같으면 ID도 같다.
런타임은 JSON만 읽으므로 openpyxl은 이 스크립트에서만 쓴다 (dev 의존성).

실행: uv run python -m app.standard.importer
"""

import json
import re
import sys
from pathlib import Path

from app.agent.lesson_prep.textbook import get_data_dir

SOURCE = {
    "title": "2017년 국제 통용 한국어 표준 교육과정 적용 연구(4단계) 어휘, 문법 등급 목록 (2020-11-17 수정)",
    "publisher": "국립국어원",
    "url": "https://www.korean.go.kr/front/reportData/reportDataView.do?mn_id=45&report_seq=932",
    "license": "공공누리 제1유형 (출처표시)",
}

# "가격02", "은1", "-는데2"처럼 동음이의 번호가 붙은 표제어를 (글자, 번호)로 나눈다
HOMONYM = re.compile(r"^(.*?)(\d+)$")


# 품사가 여러 개인 표제어는 "그02/그01", "첫째02‧첫째01/첫째03"처럼 한 칸에 묶여 있다 (206개)
HEADWORD_SEP = re.compile(r"[/‧∙·・]")


def split_homonym(text: str) -> tuple[str, str | None]:
    match = HOMONYM.match(text)
    return (match.group(1), match.group(2)) if match else (text, None)


def split_headword(text: str) -> tuple[str, str | None]:
    """'그02/그01' -> ('그', '02/01'). 묶인 표제어는 모두 같은 글자라서 글자 하나로 합친다."""
    parts = [split_homonym(p.strip()) for p in HEADWORD_SEP.split(text) if p.strip()]
    words = {word for word, _ in parts}
    if len(words) != 1:
        raise ValueError(f"묶인 표제어의 글자가 서로 다르다: {text}")
    numbers = [n for _, n in parts if n]
    return parts[0][0], "/".join(numbers) or None


def level_no(text: str) -> int:
    """'1급' -> 1"""
    return int(text.strip().rstrip("급"))


def clean(value) -> str | None:
    if value is None:
        return None
    text = str(value).strip()
    return text or None


def read_grammar(sheet) -> list[dict]:
    rows = []
    # 칸: 전체 번호, 등급별 번호, 등급, 분류, 대표형, 관련형, 의미, ...
    for row in list(sheet.iter_rows(values_only=True))[1:]:
        if row[0] is None:
            continue
        form, homonym = split_homonym(clean(row[4]))
        related = [split_homonym(r.strip())[0] for r in (clean(row[5]) or "").split(",") if r.strip()]
        rows.append(
            {
                "id": f"std-g-{int(row[0]):03d}",
                "level": level_no(row[2]),
                "category": clean(row[3]),
                "form": form,
                "homonym": homonym,
                "related": related,
                "meaning": clean(row[6]),
            }
        )
    return rows


def read_vocab(sheet) -> list[dict]:
    rows = []
    # 칸: 전체 번호, 등급별 번호, 등급, 어휘, 품사, 길잡이말, ...
    for row in list(sheet.iter_rows(values_only=True))[1:]:
        if row[0] is None:
            continue
        word, homonym = split_headword(clean(row[3]))
        rows.append(
            {
                "id": f"std-v-{int(row[0]):05d}",
                "level": level_no(row[2]),
                "word": word,
                "homonym": homonym,
                "pos": clean(row[4]),
                "guide": clean(row[5]),
            }
        )
    return rows


def import_standard(xlsx: Path, out_dir: Path) -> None:
    from openpyxl import load_workbook

    workbook = load_workbook(xlsx, read_only=True)
    for name, reader, filename in [("문법", read_grammar, "grammar.json"), ("어휘", read_vocab, "vocab.json")]:
        items = reader(workbook[name])
        (out_dir / filename).write_text(
            json.dumps({"source": SOURCE, "items": items}, ensure_ascii=False, indent=1) + "\n", encoding="utf-8"
        )
        print(f"{filename}: {len(items)}개")


if __name__ == "__main__":
    out = get_data_dir() / "standard"
    src = Path(sys.argv[1]) if len(sys.argv) > 1 else next((out / "source").glob("*.xlsx"))
    import_standard(src, out)
